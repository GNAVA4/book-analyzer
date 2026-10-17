"""Разбор DOCX-работы (курсовая/отчёт) в формат parse_article_pdf."""
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.services.article_docx import parse_article_docx

LONG = 'Предложение основного текста работы с достаточно длинным содержанием. ' * 6


def make(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run('Исследование методов анализа текстов').bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph('Выполнил: Иванов И. И., студент группы ИВТ-01')
    d.add_heading('Содержание', 1)
    d.add_paragraph('Введение ........................ 3')
    d.add_heading('Введение', 1)
    d.add_paragraph(LONG)
    d.add_paragraph('Задачи работы:')
    d.add_paragraph('– собрать корпус;', style='List Paragraph')
    d.add_paragraph('– проверить метод.', style='List Paragraph')
    d.add_heading('1.1 Метод', 2)
    d.add_paragraph(LONG)
    c = d.add_paragraph('Рисунок 1 – Схема метода')
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_heading('Список использованных источников', 1)
    d.add_paragraph('1. Петров П. П. Анализ текстов. – М., 2020.')
    path = tmp_path / 'kurs.docx'
    d.save(path)
    return str(path)


def test_structure_and_meta(tmp_path):
    r = parse_article_docx(make(tmp_path))
    assert r['meta']['title'] == 'Исследование методов анализа текстов'
    assert r['meta']['authors'] == ['Иванов И. И.']
    assert [(s['title'], s['level']) for s in r['sections']] == [('Введение', 1), ('1.1 Метод', 2)]
    assert len(r['meta']['references']) == 1
    assert r['lang'] == 'ru' and not r['needs_ocr']


def test_toc_caption_skipped_and_list_joined(tmp_path):
    r = parse_article_docx(make(tmp_path))
    intro = r['sections'][0]['paragraphs']
    assert intro[1] == 'Задачи работы:\n– собрать корпус;\n– проверить метод.'
    body = ' '.join(t for s in r['sections'] for t in s['paragraphs'])
    assert '......' not in body and 'Рисунок 1' not in body
